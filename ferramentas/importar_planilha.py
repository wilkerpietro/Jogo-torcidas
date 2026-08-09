#!/usr/bin/env python3
"""
Importa a planilha do autor para os arquivos de dados do jogo.

    python3 ferramentas/importar_planilha.py dados/fonte/Book_3_1.xlsx

Gera dados/times.js. As cidades e as torcidas tem fonte propria,
mais rica, e sao geradas por outros dois importadores:

    importar_bairros.py   -> dados/cidades.js  (le tambem esta planilha,
                             so pro efetivo de rua: guardas, PMs, choque)
    importar_relacoes.py  -> dados/torcidas.js

Rodar de novo depois de editar a planilha; nada e escrito a mao.
"""
import json, pathlib, sys, unicodedata

try:
    import openpyxl
except ImportError:
    sys.exit('falta openpyxl: pip install openpyxl')

RAIZ = pathlib.Path(__file__).resolve().parent.parent

# nomes de cor da planilha -> hex. Tons escolhidos pra ler bem sobre fundo escuro.
CORES = {
    'branco':'#e9e9e9', 'preto':'#151515', 'vermelho':'#c8102e', 'verde':'#0f7a3d',
    'azul':'#1b4f9c', 'azul claro':'#4a9fd8', 'azul marinho':'#122a52',
    'amarelo':'#e8c020', 'grena':'#6e1a2b', 'cinza':'#8a8a8a',
}
PADRAO = '#8a8a8a'

# a aba Times chama de Meio-Norte o que a aba Cidades chama de Maranhao.
# E a mesma praca (Sao Luis, Imperatriz, Teresina) — o GDD 18.3 ja trata
# Piaui e Maranhao juntos. Apelido resolve sem mexer na planilha.
APELIDOS_MAPA = {'meio-norte': 'maranhao'}


def sem_acento(t):
    return ''.join(c for c in unicodedata.normalize('NFD', t)
                   if unicodedata.category(c) != 'Mn')


def cor(nome):
    if not nome:
        return None
    chave = sem_acento(str(nome).strip().lower())
    return CORES.get(chave, PADRAO)


def ident_mapa(v):
    return APELIDOS_MAPA.get(ident(v), ident(v))


def ident(texto):
    """id estavel, sem acento nem espaco — serve de chave entre os arquivos"""
    t = sem_acento(str(texto).strip().lower())
    saida = ''.join(c if c.isalnum() else '-' for c in t)
    while '--' in saida:
        saida = saida.replace('--', '-')
    return saida.strip('-')


def texto(v):
    if v is None:
        return ''
    s = str(v).strip()
    return '' if s == '-' else s


def inteiro(v, padrao=0):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return padrao


def escrever(caminho, cabecalho, nome_campo, itens):
    linhas = ',\n'.join('  ' + json.dumps(i, ensure_ascii=False) for i in itens)
    (RAIZ / caminho).write_text(
        f'/* {cabecalho}\n'
        f'   GERADO por ferramentas/importar_planilha.py — nao editar a mao.\n'
        f'   {len(itens)} registros. */\n'
        f'{nome_campo} = [\n{linhas}\n];\n', encoding='utf-8')
    print(f'  {caminho}: {len(itens)} registros')


def main():
    origem = sys.argv[1] if len(sys.argv) > 1 else 'dados/fonte/Book_3_1.xlsx'
    wb = openpyxl.load_workbook(RAIZ / origem, data_only=True)

    # ---------------- cidades ----------------
    # o mapa sai dos quarteiroes; o resto e efetivo de rua e de policia
    cidades = []
    for r in list(wb['Níveis de cidades'].iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        quart = inteiro(r[7], 64)
        lado = {144: 12, 100: 10, 64: 8}.get(quart, 8)
        cidades.append({
            'id': ident(r[0]), 'nome': texto(r[0]),
            'torcedores': inteiro(r[1]),
            'nivel': inteiro(texto(r[2]).replace('Nível', '').strip(), 3),
            'guardas': inteiro(r[3]), 'pms': inteiro(r[4]), 'choque': inteiro(r[5]),
            'personagens': inteiro(r[6]),
            'quarteiroes': quart, 'grade': [lado, lado],
            'npcsPorQuarteirao': round(float(r[8] or 0), 2),
        })

    # ---------------- times ----------------
    times = []
    for r in list(wb['Times'].iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        cores = [c for c in (cor(r[8]), cor(r[9]),
                             cor(r[11]) if texto(r[10]).lower() == 'sim' else None) if c]
        times.append({
            'id': ident(r[0]), 'nome': texto(r[0]), 'nomeCompleto': texto(r[1]),
            'sigla': texto(r[2]), 'alcunha': texto(r[3]),
            'cidade': texto(r[4]), 'uf': texto(r[5]),
            'estadio': texto(r[6]), 'capacidade': inteiro(r[7]),
            'cores': cores or [PADRAO],
            'fundacao': inteiro(r[12]), 'mascote': texto(r[13]),
            'qualidade': inteiro(r[14]),
            'divisao': texto(r[15]), 'mapa': ident_mapa(r[16]),
            'regional': texto(r[17]),
        })

    # ---------------- torcidas ----------------
    por_time = {t['nome'].lower(): t for t in times}
    torcidas, sem_time = [], []
    for r in list(wb['Torcidas'].iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        nome_time = texto(r[2])
        t = por_time.get(nome_time.lower())
        if not t:
            sem_time.append(nome_time)
        fund = r[1]
        torcidas.append({
            'id': ident(r[0]), 'nome': texto(r[0]),
            'fundacao': fund.year if hasattr(fund, 'year') else inteiro(fund),
            'time': t['id'] if t else ident(nome_time),
            'timeNome': nome_time,
            'mapa': ident_mapa(r[3]),
            'cores': [cor(r[4]) or PADRAO, cor(r[5]) or PADRAO],
            'linha': cor(r[6]) or PADRAO,
            'membros': inteiro(r[7]),
        })

    # cidades.js NAO e escrito aqui. Quem manda nele e
    # ferramentas/importar_bairros.py, que junta o JSON de bairros com o
    # efetivo de rua desta planilha. Escrever aqui apagaria os 348 bairros.
    escrever('dados/times.js', 'TIMES — clubes, estadios e divisoes',
             'TO.dados.times', times)
    # torcidas.js tambem nao: vem de ferramentas/importar_relacoes.py,
    # que tem bairro-sede e o grafo de relacoes.

    # ---------------- conferencia ----------------
    print('\nconferencia:')
    mapas = {c['id'] for c in cidades}
    orfaos_t = sorted({t['mapa'] for t in times if t['mapa'] not in mapas})
    orfaos_o = sorted({o['mapa'] for o in torcidas if o['mapa'] not in mapas})
    print(f"  cidades por grade: " + ', '.join(
        f"{lado}x{lado}={sum(1 for c in cidades if c['grade'][0]==lado)}"
        for lado in (12, 10, 8)))
    print(f'  times sem cidade conhecida: {orfaos_t or "nenhum"}')
    print(f'  torcidas sem cidade conhecida: {orfaos_o or "nenhum"}')
    if sem_time:
        print(f'  torcidas cujo time nao esta na aba Times: {sorted(set(sem_time))}')
    else:
        print('  todas as torcidas apontam pra um time existente')
    semt = [c['nome'] for c in cidades
            if not any(t['mapa'] == c['id'] for t in times)]
    print(f'  cidades sem nenhum time: {semt or "nenhuma"}')


if __name__ == '__main__':
    main()
