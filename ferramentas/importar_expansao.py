#!/usr/bin/env python3
"""
Junta a planilha do Brasil com a da expansao sul-americana e escreve
dados/times.js inteiro.

    python3 ferramentas/importar_expansao.py

Le dados/fonte/Book_3_1.xlsx (aba Times, 108 clubes do Brasil) e
dados/fonte/Expansao_America_do_Sul.xlsx (aba Times, 248 clubes de nove
paises). Nada e escrito a mao; rodar de novo depois de editar qualquer
uma das duas.

DUAS COLUNAS NOVAS no registro do clube:
  pais  — 'Brasil' ou o pais da expansao. E o que separa o Brasileirao
          das ligas de fora e mantem a Copa do Brasil so com brasileiro.
  copa  — a copa nacional do clube. No Brasil ela sai do proprio motor
          (Copa do Brasil); nos outros vem da planilha.
A coluna `regional` continua sendo o estadual, e clube de fora nao tem.
"""
import json, pathlib, sys, unicodedata

try:
    import openpyxl
except ImportError:
    sys.exit('falta openpyxl: pip install openpyxl')

RAIZ = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / 'ferramentas'))
from importar_planilha import CORES, PADRAO, cor, ident, ident_mapa, texto, \
                              inteiro, escrever, REGIONAL_CORRIGIDO


def do_brasil(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    fora = []
    for r in list(wb['Times'].iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        cores = [c for c in (cor(r[8]), cor(r[9]),
                             cor(r[11]) if texto(r[10]).lower() == 'sim' else None) if c]
        fora.append({
            'id': ident(r[0]), 'nome': texto(r[0]), 'nomeCompleto': texto(r[1]),
            'sigla': texto(r[2]), 'alcunha': texto(r[3]),
            'cidade': texto(r[4]), 'uf': texto(r[5]),
            'estadio': texto(r[6]), 'capacidade': inteiro(r[7]),
            'cores': cores or [PADRAO],
            'fundacao': inteiro(r[12]), 'mascote': texto(r[13]),
            'qualidade': inteiro(r[14]),
            'divisao': texto(r[15]), 'mapa': ident_mapa(r[16]),
            'regional': REGIONAL_CORRIGIDO.get(ident(r[0]), texto(r[17])),
            'pais': 'Brasil',
        })
    return fora


def da_expansao(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    fora = []
    for r in list(wb['Times'].iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        cores = [c for c in (cor(r[8]), cor(r[9]),
                             cor(r[11]) if texto(r[10]).lower() == 'sim' else None) if c]
        fora.append({
            'id': ident(r[0]), 'nome': texto(r[0]), 'nomeCompleto': texto(r[1]),
            'sigla': texto(r[2]), 'alcunha': texto(r[3]),
            'cidade': texto(r[4]), 'uf': texto(r[5]),
            'estadio': texto(r[6]), 'capacidade': inteiro(r[7]),
            'cores': cores or [PADRAO],
            'fundacao': inteiro(r[12]), 'mascote': texto(r[13]),
            'qualidade': inteiro(r[14]),
            'divisao': texto(r[15]), 'mapa': ident(r[16]),
            'regional': '',                 # la fora nao ha estadual
            'copa': texto(r[17]),           # a copa nacional do pais
            'pais': texto(r[18]),
        })
    return fora


# A TROCA DE C E D (decisao do dono, 21/08/2026). A planilha do autor
# nunca foi atualizada, entao a decisao vivia so no times.js gerado --
# e a importacao da expansao a apagou junto. Agora ela mora aqui, que e
# onde sobrevive a proxima regeracao. DECISOES.md guarda o porque.
TROCA_C_D = {
    'amazonas':    ('Brasileirão Série D', 6),
    'ferroviario': ('Brasileirão Série D', 6),
    'floresta':    ('Brasileirão Série D', 7),
    'brasiliense': ('Brasileirão Série C', 10),
    'sergipe':     ('Brasileirão Série C', 10),
    'csa':         ('Brasileirão Série C', 8),
    'joinville':        ('Brasileirão Série C', 8),
    'inter-de-limeira': ('Brasileirão Série D', 6),
}


def aplicar_remendos(times):
    achados = 0
    for t in times:
        r = TROCA_C_D.get(t['id'])
        if r:
            t['divisao'], t['qualidade'] = r
            achados += 1
    faltando = sorted(set(TROCA_C_D) - {t['id'] for t in times})
    if faltando:
        sys.exit(f'REMENDO SEM CLUBE, a decisao do dono se perde: {faltando}')
    return achados


def main():
    br = do_brasil(RAIZ / 'dados/fonte/Book_3_1.xlsx')
    ex = da_expansao(RAIZ / 'dados/fonte/Expansao_America_do_Sul.xlsx')
    times = br + ex

    vistos, repetidos = set(), []
    for t in times:
        if t['id'] in vistos:
            repetidos.append(t['id'])
        vistos.add(t['id'])
    if repetidos:
        sys.exit(f'ID REPETIDO, o jogo funde os dois clubes: {repetidos}')

    remendados = aplicar_remendos(times)

    escrever('dados/times.js', 'TIMES — clubes, estadios e divisoes',
             'TO.dados.times', times)

    print('\nconferencia:')
    print(f'  Brasil: {len(br)} · expansao: {len(ex)} · total: {len(times)}')
    print(f'  remendos da decisao do dono aplicados: {remendados}')
    porpais = {}
    for t in times:
        porpais[t['pais']] = porpais.get(t['pais'], 0) + 1
    for p, n in sorted(porpais.items(), key=lambda x: -x[1]):
        print(f'    {p}: {n}')
    divs = {}
    for t in times:
        divs.setdefault(t['divisao'], 0)
        divs[t['divisao']] += 1
    print(f'  divisoes: {len(divs)}')
    copas = sorted({t.get('copa') for t in times if t.get('copa')})
    print(f'  copas nacionais: {len(copas)} · {", ".join(copas)}')


if __name__ == '__main__':
    main()
